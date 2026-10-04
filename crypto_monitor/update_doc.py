from docx import Document
doc = Document("策略版本记录.docx")
# 加新段落
doc.add_heading('v2.1.1 (2026-09-26 21:35)', level=2)
doc.add_paragraph('1. 新增双周期共振规则：15m和1h同时出买入/卖出信号才推送，单周期信号不推，减少噪音')
doc.add_paragraph('2. 修复重大bug：space变量未定义导致推送表格后续内容全部丢失')
doc.add_paragraph('3. 新增功能：到价提醒（BSV22U止盈/21U止损）、仓位建议（日线向上3成/向下1成）、回测算0.1%双边手续费')
doc.add_paragraph('4. 新增一键检查所有币状态脚本check_all_alerts.py')
doc.save("策略版本记录.docx")
print("写完了")
